import json
import os
import re
import textwrap

import requests
import streamlit as st
from pptx import Presentation
from pptx.util import Pt

MODEL = "claude-haiku-4-5-20251001"
WIKI_API = "https://en.wikipedia.org/w/api.php"
HEADERS = {"User-Agent": "PPT-Generator/1.0 (student project)"}
SKIP_SECTIONS = {"see also", "references", "external links", "further reading",
                 "notes", "bibliography", "sources", "citations"}


# ---------- Option 1: free mode (Wikipedia, no API key needed) ----------
def get_wikipedia_text(topic):
    search = requests.get(WIKI_API, headers=HEADERS, timeout=15, params={
        "action": "query", "list": "search", "srsearch": topic,
        "srlimit": 1, "format": "json"}).json()
    results = search["query"]["search"]
    if not results:
        raise ValueError(f"No Wikipedia article found for '{topic}'. Try another topic.")
    title = results[0]["title"]

    data = requests.get(WIKI_API, headers=HEADERS, timeout=15, params={
        "action": "query", "prop": "extracts", "explaintext": 1,
        "titles": title, "redirects": 1, "format": "json"}).json()
    page = next(iter(data["query"]["pages"].values()))
    return page["extract"]


def good_sentences(text, limit):
    sentences = re.split(r"(?<=[.!?])\s+", text.replace("\n", " "))
    bullets = []
    for sentence in sentences:
        sentence = sentence.strip()
        if len(sentence) >= 30:
            bullets.append(textwrap.shorten(sentence, width=150, placeholder="..."))
        if len(bullets) == limit:
            break
    return bullets


def slides_from_text(text, content_slides):
    parts = re.split(r"\n=+ (.+?) =+\n", text)
    # parts = [intro, title1, body1, title2, body2, ...]
    sections = [("Introduction", parts[0])]
    for i in range(1, len(parts) - 1, 2):
        if parts[i].strip().lower() not in SKIP_SECTIONS:
            sections.append((parts[i].strip(), parts[i + 1]))

    slides = []
    for title, body in sections:
        bullets = good_sentences(body, 3)
        if bullets:
            slides.append({"title": title, "bullets": bullets})
        if len(slides) == content_slides:
            break
    return slides


def generate_slides_free(topic, number_of_slides):
    text = get_wikipedia_text(topic)
    slides = slides_from_text(text, number_of_slides - 1)
    if not slides:
        raise ValueError("Could not find enough content for this topic.")
    return slides


# ---------- Option 2: AI mode (Claude, needs API key) ----------
def generate_slides_ai(topic, number_of_slides, api_key):
    import anthropic

    client = anthropic.Anthropic(api_key=api_key)
    content_slides = number_of_slides - 1

    prompt = f"""Create content for a presentation on "{topic}".
Write exactly {content_slides} content slides in a logical order
(for example: introduction, key concepts, features, applications,
advantages, challenges, future scope, conclusion).

Return ONLY valid JSON, with no extra text and no markdown fences.
Use this format:
[
  {{"title": "Short slide title", "bullets": ["point 1", "point 2", "point 3", "point 4"]}}
]
Each slide must have 4 bullets, and each bullet must be under 15 words."""

    response = client.messages.create(
        model=MODEL,
        max_tokens=3000,
        messages=[{"role": "user", "content": prompt}],
    )
    text = response.content[0].text.strip()
    text = text.replace("```json", "").replace("```", "").strip()
    return json.loads(text)


# ---------- Build the PowerPoint ----------
def build_presentation(topic, slides, file_name, subtitle):
    presentation = Presentation()

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = topic.title()
    title_slide.placeholders[1].text = subtitle

    for item in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = item["title"]

        body = slide.placeholders[1].text_frame
        body.clear()
        for index, bullet in enumerate(item["bullets"]):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = bullet
            paragraph.font.size = Pt(20)

    presentation.save(file_name)


# ---------------- Streamlit app ----------------
st.set_page_config(page_title="PPT Generator", page_icon="🤖", layout="centered")

st.title("🤖 PPT Presentation Generator")
st.write("Enter a topic and get a PowerPoint with real content on every slide.")

with st.sidebar:
    st.header("Settings")
    typed_key = st.text_input("Anthropic API key (optional)", type="password")
    st.caption("Leave empty to use the free mode (Wikipedia). "
               "Add a key to let AI write the slides.")

api_key = typed_key or os.environ.get("ANTHROPIC_API_KEY", "")

topic = st.text_input(
    "Enter your presentation topic",
    placeholder="Example: Cloud Computing",
)
number_of_slides = st.slider("Number of slides", min_value=3, max_value=10, value=5)

if st.button("Generate PPT 🚀"):
    if topic.strip() == "":
        st.warning("Please enter a topic first.")
    else:
        try:
            with st.spinner("Writing your slides..."):
                if api_key:
                    slides = generate_slides_ai(topic, number_of_slides, api_key)
                    subtitle = "Generated using Streamlit and AI"
                else:
                    slides = generate_slides_free(topic, number_of_slides)
                    subtitle = "Generated using Streamlit and Wikipedia"
                file_name = "Presentation.pptx"
                build_presentation(topic, slides, file_name, subtitle)

            st.success(f"🎉 Presentation created with {len(slides) + 1} slides!")

            with open(file_name, "rb") as file:
                st.download_button(
                    label="⬇️ Download PPT",
                    data=file,
                    file_name=file_name,
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
        except Exception as error:
            st.error(f"Something went wrong: {error}")